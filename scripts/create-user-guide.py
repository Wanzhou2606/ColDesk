from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output'/'user-guide'
OUT.mkdir(parents=True,exist_ok=True)
IMAGES=ROOT/'preview'/'goals-guide'
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
section.left_margin=Inches(.8);section.right_margin=Inches(.8)
section.header_distance=Inches(.25);section.footer_distance=Inches(.3)
for name,size in [('Normal',11),('Title',28),('Subtitle',12),('Heading 1',21),('Heading 2',14),('Caption',9)]:
    style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    style.font.bold=name.startswith('Heading')
    style.font.italic=False
    style.paragraph_format.line_spacing=Pt(16.5 if name=='Normal' else 13 if name=='Caption' else size*1.3)
    style.paragraph_format.space_after=Pt(7)
    if name.startswith('Heading'):style.paragraph_format.space_before=Pt(12 if name=='Heading 2' else 0)
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(snap)
doc.styles['Normal'].paragraph_format.keep_with_next=False
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('637B82')
doc.styles['Caption'].paragraph_format.space_after=Pt(12)
doc.styles['Normal'].paragraph_format.widow_control=True
header=section.header.paragraphs[0];header.text='ColDesk 用户手册';header.style=doc.styles['Caption'];header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=section.footer.paragraphs[0];footer.alignment=2
run=footer.add_run('ColDesk 0.8.0   ');run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='ColDesk 用户手册'
doc.core_properties.subject='Windows 版日程 日记 长期任务与个人设置使用说明'
doc.core_properties.author='ColDesk'

def p(text='',bold_lead=None):
    paragraph=doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):paragraph.add_run(bold_lead).bold=True;paragraph.add_run(text[len(bold_lead):])
    else:paragraph.add_run(text)
    return paragraph
def h(text):doc.add_heading(text,2)
def page(text):doc.add_page_break();doc.add_heading(text,1)
def steps(items):
    for number,text in enumerate(items,1):
        paragraph=p(f'{number}. {text}');paragraph.paragraph_format.left_indent=Inches(.17);paragraph.paragraph_format.first_line_indent=Inches(-.17)
def picture(name,caption,width=6.9,crop=None):
    source=IMAGES/(name+'.png')
    with Image.open(source) as image:w,h=image.size
    paragraph=doc.add_paragraph();paragraph.paragraph_format.space_after=Pt(3)
    paragraph.paragraph_format.line_spacing=1.0
    paragraph.paragraph_format.keep_with_next=True
    shape=paragraph.add_run().add_picture(str(source),width=Inches(width))
    if crop:
        left,top,right,bottom=crop
        source_rect=OxmlElement('a:srcRect')
        for key,value in zip(['l','t','r','b'],crop):source_rect.set(key,str(int(value*1000)))
        fill=shape._inline.graphic.graphicData.pic.blipFill
        fill.insert(1,source_rect)
        shape.height=Inches(width*(h*(100-top-bottom))/(w*(100-left-right)))
    doc.add_paragraph(caption,'Caption')

doc.add_paragraph('ColDesk 用户手册','Title')
doc.add_paragraph('Windows 版 0.8.0','Subtitle')
p('本手册面向第一次使用 ColDesk 的用户，介绍日程、任务、长期方向、日记、打卡、记账与个性化设置。你可以先添加一项任务，再把它拖进日程；需要记录心情或长期进展时，再进入 Home 和长期页面。')
picture('schedule-day','图 1  日程是软件默认打开的页面。本文截图均使用独立的演示数据。',width=6.9)
h('第一次打开')
steps(['双击桌面上的 ColDesk 图标，或在 Windows 开始菜单搜索 ColDesk。首次使用按引导完成基本设置，也可以先跳过。','启动汇总会列出今天的课程与紧急任务。查看完后点击“收到”；即使 Windows 通知收起，应用内确认卡片仍会保留。','左侧默认顺序是 Home、日程、任务、打卡、长期、记账。语言、导航顺序和启动页面都能在设置的通用栏目中修改。'])
p('需要再看教程时，进入“设置 → 使用帮助 → 重新查看引导”。在 Home 中先点左下角三横杠菜单，再进入设置。')

page('日程与课程视图')
p('日程提供日程表、清单和周视图。日程表适合安排当天时间，清单适合按时间顺序浏览，周视图适合查看整周课程与安排。')
picture('schedule-week','图 2  周视图中的课程与任务会一起显示。课表模式仅改变时间轴刻度。',width=6.9)
steps(['在右侧“近期截止”中填写本次时长，把任务拖到日程表的空闲时间。与课程或固定安排冲突时不能放置。','已经放入日程的任务块也可以拖动。再次拖入同一任务会移动原安排，不会重复创建。','用“＋”“−”调整缩放，用“总览”恢复预览。日程默认显示 06:00 到 24:00，可以在“设置 → 课程与排程 → 日程时间”中修改。','点击周视图中的课程卡片，查看老师、教室、实际起止时刻和节数。点击空白处或按 Esc 收起；其他课程保持原位。'])
p('打开“课表模式”后，日程表和周视图只显示每节的开始时间、数字节次与结束时间，不混入整点刻度。通用订阅缺少节次信息时，会标注估算或按实际时段编号。红线会随当前时间移动。')

page('导入课程表')
p('杭电用户可以直接读取官方课表订阅；其他课表可以由你常用的 AI 识别图片，再导入结果。两种方式都先预览、核对，再确认。')
picture('settings-courses','图 3  课程与排程设置集中放置日程时间、官方订阅和图片识别导入。',width=6.9,crop=(13,0,0,22))
h('杭电助手订阅')
steps(['点击日程右上角“导入课表”，或进入“设置 → 课程与排程”。','打开杭电助手并登录，在官方课表右上角菜单选择“订阅课表”，复制链接。回到 ColDesk 粘贴并读取。','核对课程日期、时刻、老师与教室，确认后使用这份课表。链接仅用于本次读取；课程变化后，需要重新粘贴读取更新。'])
h('课表图片识别')
p('在同一设置页复制识别提示词，把提示词和课表图片交给你常用的 AI，再把输出粘回导入框。常见的字段、星期和周次写法会自动整理；格式大致相同即可。缺少学期起始日期或节次时刻时，需要提供相应信息。')
p('结果仅供参考，请仔细核对。重点检查科目、星期、节次时间、上课周次和教室；确认替换课表后，原有任务、日记、打卡和记账记录仍保留。',bold_lead='结果仅供参考，请仔细核对。')

page('任务与桌面组件')
p('任务页包含学习任务和生活任务两个标签。任务负责记录要做的事，日程负责给任务安排具体时间。')
picture('tasks','图 4  学习任务按期限分组。生活事项保留在同一页面的另一标签中。',width=6.9,crop=(13,0,0,15))
steps(['点击“新建任务”，填写名称、预计用时、紧迫程度和截止时间。生活任务也可以在顶部输入框随手记录，按回车保存。','任务右上角的“⋯”可以删除、置顶、修改紧迫程度、截止时间和名称。删除前会确认，并清理相关日程和提醒。','自动安排开启后，会按可用时间、截止时间等排程并避开课程。学习和生活的可用时段及单日上限可以分别设置；也能关闭自动安排后手动拖入日程。'])
h('桌面悬浮窗与灵动岛')
p('在“设置 → 外观与桌面 → 桌面形态”中选择悬浮窗、贴边抽屉或灵动岛。拖住顶部移动，松手保存位置；灵动岛在悬停时展开、移开时收起，提醒时动态弹开。置顶、普通与置底控制它是否会被其他窗口遮住。')
p('Ctrl + Alt + T 用于快速添加任务；Ctrl + Alt + D 显示或隐藏悬浮窗。关闭主窗口后软件仍在托盘运行，右键托盘图标可以重新打开或退出。')

page('建立一个长期方向')
p('长期用于“学习日语”“完成一个作品”等较大的方向。先建立一个方向，不需要一开始就知道共有多少步骤。它独立于日程任务，不会自动把整个方向塞进某个时间段。')
picture('goals-list','图 5  先看到长期任务列表，再选择要继续推进的方向。',width=6.9,crop=(14,0,0,45))
steps(['进入“长期”，点击“新建长期任务”。填写名称和说明，例如“学习日语”，保存。','在列表中点击一个方向，卡片会展开为方向手册。点击左上角“长期任务列表”可以回到列表，选择其他方向。','点击“添加分支”写下一个目前想推进的方面，例如“五十音与发音”“基础词汇与语法”。以后想到新的内容再添加。'])
h('理解页面中的数量')
p('“当前记录”只统计已经添加的项目，不代表整个方向的最终总步骤。“已完成”表示当前已标记完成的分支数量。软件不会因为现有分支都完成，就自动宣告整个长期方向结束。')
p('编辑按钮可以修改名称和说明。删除一个长期任务会同时删除它的下级分支及关联信息；已经关联的日记原文会保留。删除前请确认。')

page('逐级细分与完成记录')
picture('goals-handbook','图 6  方向手册在左侧展开多级分支，右侧查看日记足迹与最近完成。',width=6.9,crop=(14,0,0,8))
steps(['点击任意分支右侧的“＋”，继续添加下级分支。分支可以一级一级细分，没有预设层数和数量上限；箭头用于展开或收起下级。','点圆圈标记完成，完成时间会写入数据库。鼠标停留在项目上，或用键盘聚焦其按钮，可看到完成时间，其他项目保持原位。父项目需要先完成下级，才能标记完成。','重新打开已完成项目时，原完成时间保留在历史中。给已完成分支添加新下级，也会重新打开相关上级；之后再次完成会记录新的时间。','点击分支名称，在右侧选择日期并添加日记关联。从足迹点击日期，即可前往该天日记。加密日记仍需密码解锁，解除关联不会删除日记。'])

page('Home 日记与月历')
p('Home 用来记录一天的故事。第一次进入时，可以设置密码，也可以选择不设置；之后可在 Home 右上角调整密码保护。')
picture('home-calendar','图 7  月视图显示每天的封面文字，点击日期进入日记。',width=6.9)
steps(['点击某一天，在最上方“封面”中写一句话。返回月历后，这句话会以手写体显示在当天格子里。','选择天气和心情，填写小标题与正文。小标题会用于日记导出的标题；封面与正文可以分别记录。','文字会自动保存，返回月历或离开 Home 前也会先保存。若出现保存失败提示，请重试后再离开。','Home 会把左侧导航收进左下角三横杠菜单。点菜单即可切换栏目或进入设置；离开有密码保护的 Home 后，下次进入需要重新解锁。'])
p('请记住日记密码。忘记密码后，软件无法恢复加密日记。文字、照片和文件都受同一保护规则约束。')

page('照片 文件 导出与关联')
picture('home-export','图 8  导出可选择日期范围、按天保存或合成总览。',width=6.9)
h('把日记作为相册')
p('在日记编辑页下方用“＋ 照片”和“＋ 文件”添加附件。切换相册可按月浏览照片，点照片查看大图、切换前后照片或导出原图。桌面版单个附件最多 25 MB，每批最多 20 个、合计 100 MB，每天最多 60 个。导入后使用本地副本，移动原文件不会改变已保存附件。')
h('导出自己的记录')
steps(['在 Home 右上角点击“导出日记”，选择当天、当前月份、全部或自选起止日期。','选择“总览”生成一份包含文字和照片的 Word 文档；选择“日记”则按天生成文件，标题使用日期与当天小标题。','无照片时可以选 TXT；含照片时自动使用 Word。其他文件保存到相邻的附件目录。选择目标文件夹后开始导出。'])
p('加密日记需要先解锁再导出。导出的 Word、TXT 和原图不再受 Home 密码保护，请妥善保管；原日记仍会留在软件中。')
h('日记与记账双向关联')
p('在日记编辑页的“关联记账”中选择一笔流水并添加；也可以在记账流水旁点击“关联日记”，选择日期。保存后两边都能点击跳转。长期页面同样可以关联日记日期，用来记录方向与分支的进展。')

page('自律打卡')
picture('habits','图 9  今日打卡、习惯管理和月度记录集中在同一页面。',width=6.9,crop=(13,0,0,12))
steps(['点击“新建习惯”，输入一项想每天坚持的事。新习惯默认每天打卡。','完成后点“今日打卡”中的相应项目；再次点击可撤销误操作。月历会汇总每天的完成情况。','点击“管理”可改名、归档或删除习惯。暂时不做时可以先归档；删除会一并删除该习惯的签到记录。'])
h('回看记录与每日名言')
p('月历可以切换月份，查看过去的完成情况。每天不必追求所有事情都做满，先保持适合自己的习惯数量。页面上的名言可以“换一句”或复制；网络不可用时仍会显示内置内容。')
p('打卡记录不会自动完成长期分支，也不会自动完成日程中的任务。如果同一件事同时出现在这些地方，需要在对应栏目分别记录。')

page('记账与预算')
picture('expenses','图 10  记账页包含区间统计、分类和流水。',width=6.9,crop=(13,0,0,12))
steps(['选择分类，输入大于 0 的金额，按“记下”保存一笔支出。备注可用于区分用途。这个模块只记录支出。','用日、周、月切换查看统计。需要预算提示时设置每日总预算和分类目标；不设置预算也可以正常记账。','每月可以填写期初可用余额。填写后优先使用它作为当月预算，未填写时按每日总预算换算。','流水旁可以关联日记日期，点击关联即可跳到那一天。删错流水时请谨慎确认；解除关联只取消链接，删除流水也会清理相应链接。'])
p('分类名称、金额、备注和记录都是你自己的数据，切换界面语言不会翻译或改写这些内容。预算提示只是记录与统计，不会替你移动资金。')

page('图标 主题与字体')
picture('settings-theme','图 11  个性化中可以分别选择主页面与 Home 的颜色，包括新增的雾青。',width=6.9,crop=(13,0,0,12))
h('换成喜欢的样子')
p('进入“设置 → 个性化”，上传图片作为软件图标，或恢复默认。PNG、JPG、WebP 等图片最大 20 MB，会居中裁成方形，并保留透明区域。Windows 的快捷方式图标缓存可能需要稍等才更新。')
p('主页面和 Home 的主题颜色可以分别选择。雾青沿用方向手册方案里的浅灰蓝背景、柔和青绿和暖白卡片；也可选默认、浅蓝、浅绿、淡紫或粉色。设置立即生效，重启后保留。')
h('字体与可读性')
p('在“设置 → 外观与桌面 → 字体与月历”中，分别设置应用字体和 Home 月历封面字体。可以选择预设，也可以填写本机已安装的字体名称。未安装的字体会回退到默认；提高对比度可让文字和边界更清楚。')
p('自定义图标保存在这台电脑上，不会带进给别人分享的安装包。主页面配色与 Home 配色也相互独立。')

page('启动 导航与数据保护')
picture('settings-general','图 12  通用设置支持语言、启动页面和导航顺序。',width=6.9,crop=(13,0,0,25))
p('界面支持简体中文、English 和日本語。导航顺序用上下按钮调整；启动页面默认日程，也可以改成长期、Home 或其他已提供的页面。勾选开机自启后，登录 Windows 会启动 ColDesk。')
h('关闭窗口与退出软件')
p('关闭主窗口会保留托盘运行，提醒和桌面组件仍可使用。要完整退出，请右键 Windows 右下角的 ColDesk 托盘图标选择“退出”。再次查看引导的入口是“设置 → 使用帮助 → 重新查看引导”。')
h('保存与备份')
p('个人记录保存在本机。软件启动时自动备份数据库，保留最近 14 份；在“设置 → 数据”中可打开备份文件夹。日记与附件也需要妥善保管，迁移电脑前应先正常退出软件并完整复制数据目录，不能只复制安装包。不要覆盖正在使用中的数据库。')
h('常见情况')
p('任务没有排进日程时，先检查是否开启自动安排、预计用时、截止时间和可用时段，再检查课程冲突。课表读取失败时检查网络与订阅链接；导入后的变化需要重新读取更新。找不到主窗口时，先查看系统托盘。')
p('日记保存失败时，请留在当前页面重试。不要强制结束进程；忘记加密密码时无法恢复原内容。导出文件不加密，分享安装包也不会包含个人记录。')

destination=OUT/'ColDesk用户手册.docx'
for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(destination)
print(destination)
